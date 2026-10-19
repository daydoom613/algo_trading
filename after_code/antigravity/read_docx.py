import docx
import os

def read_docx(file_path, output_file):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    try:
        doc = docx.Document(file_path)
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(f"Total paragraphs: {len(doc.paragraphs)}\n")
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                     f.write(f"--- Paragraph {i} ---\n")
                     f.write(para.text + "\n")
                     f.write("---------------------\n")
        print(f"Content written to {output_file}")
    except Exception as e:
        print(f"Error reading docx: {e}")

file_path = "AI Trading Bot_ Training, Checkpoints, Backtesting.docx"
output_file = "requirements_extracted.txt"
read_docx(file_path, output_file)
